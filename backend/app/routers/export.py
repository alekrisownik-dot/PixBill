import io
from datetime import date

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from ..db import get_db
from ..models import Expense

router = APIRouter(prefix="/api/export", tags=["export"])

SOURCES = {"manual": "ręcznie", "receipt": "paragon", "recurring": "opłata stała"}
MONEY = '#,##0.00 "zł"'


def add_sheet(wb: Workbook, title: str, header: list[str], rows: list[list], widths: list[int], money_cols: set[int]):
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = width
        for cells in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            if col in money_cols:
                cells[0].number_format = MONEY
            elif isinstance(cells[0].value, date):
                cells[0].number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


@router.get("/xlsx")
def export_xlsx(date_from: date | None = None, date_to: date | None = None, db: Session = Depends(get_db)):
    query = select(Expense).options(selectinload(Expense.items), selectinload(Expense.category))
    if date_from:
        query = query.where(Expense.spent_on >= date_from)
    if date_to:
        query = query.where(Expense.spent_on <= date_to)
    expenses = db.scalars(query.order_by(Expense.spent_on, Expense.id)).all()

    wb = Workbook()
    wb.remove(wb.active)
    add_sheet(
        wb, "Wydatki",
        ["Data", "Kwota", "Sekcja", "Sklep", "Opis", "Źródło", "Pozycji", "ID"],
        [
            [e.spent_on, e.amount_grosze / 100, e.category.name, e.merchant or "", e.description,
             SOURCES.get(e.source, e.source), len(e.items), e.id]
            for e in expenses
        ],
        [12, 12, 20, 22, 36, 14, 9, 7], {2},
    )
    add_sheet(
        wb, "Pozycje",
        ["Data", "Produkt", "Cena", "Sklep", "Sekcja", "ID wydatku"],
        [
            [e.spent_on, i.name, i.price_grosze / 100, e.merchant or "", e.category.name, e.id]
            for e in expenses for i in e.items
        ],
        [12, 36, 12, 22, 20, 11], {3},
    )
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"pixbill-{date_from or 'start'}-{date_to or date.today()}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )
