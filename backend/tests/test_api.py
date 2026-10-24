import io
from datetime import datetime, timedelta, timezone

from PIL import Image

from app.ocr import ollama


def jpeg_bytes() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (40, 80), "white").save(buf, format="JPEG")
    return buf.getvalue()


def first_category(client) -> int:
    return client.get("/api/categories").json()[0]["id"]


def test_requires_login(client):
    client.post("/api/logout")
    assert client.get("/api/categories").status_code == 401
    assert client.post("/api/login", json={"password": "zle"}).status_code == 401


def test_default_categories_seeded(client):
    names = [c["name"] for c in client.get("/api/categories").json()]
    assert "Spożywcze" in names and "Inne" in names


def test_manual_expense_and_month_stats(client):
    cats = client.get("/api/categories").json()
    food, fuel = cats[0]["id"], cats[2]["id"]
    for amount, day, cat in [(10000, "2026-09-10", food), (15000, "2026-10-02", food), (5000, "2026-10-05", fuel)]:
        r = client.post(
            "/api/expenses",
            json={"amount_grosze": amount, "spent_on": day, "category_id": cat, "description": "test"},
        )
        assert r.status_code == 201

    assert len(client.get("/api/expenses", params={"month": "2026-10"}).json()) == 2

    summary = client.get("/api/stats/month", params={"month": "2026-10"}).json()
    assert summary["total_grosze"] == 20000
    assert summary["previous_total_grosze"] == 10000
    assert summary["change_percent"] == 100.0
    food_stat = next(c for c in summary["categories"] if c["category_id"] == food)
    assert food_stat["change_percent"] == 50.0
    fuel_stat = next(c for c in summary["categories"] if c["category_id"] == fuel)
    assert fuel_stat["change_percent"] is None

    trend = client.get("/api/stats/trend", params={"months": 3, "until": "2026-10"}).json()
    assert [p["total_grosze"] for p in trend] == [0, 10000, 20000]


def test_expense_rejects_unknown_category(client):
    r = client.post("/api/expenses", json={"amount_grosze": 100, "spent_on": "2026-10-01", "category_id": 999})
    assert r.status_code == 422


def test_receipt_flow_and_photo_cleanup(client, monkeypatch):
    monkeypatch.setattr(
        ollama,
        "read_receipt",
        lambda path: ({"merchant": "Biedronka", "date": "2026-10-20", "total_grosze": 4599, "items": []}, "SUMA PLN 45,99"),
    )
    r = client.post("/api/receipts", files={"file": ("p.jpg", jpeg_bytes(), "image/jpeg")})
    assert r.status_code == 201
    receipt_id = r.json()["id"]

    receipt = client.get(f"/api/receipts/{receipt_id}").json()
    assert receipt["status"] == "processed"
    assert receipt["ocr_result"]["total_grosze"] == 4599
    assert receipt["photo_delete_after"] is None
    assert client.get(f"/api/receipts/{receipt_id}/photo").status_code == 200

    r = client.post(
        f"/api/receipts/{receipt_id}/confirm",
        json={"amount_grosze": 4599, "spent_on": "2026-10-20", "category_id": first_category(client), "merchant": "Biedronka"},
    )
    assert r.status_code == 201
    assert r.json()["source"] == "receipt"

    receipt = client.get(f"/api/receipts/{receipt_id}").json()
    assert receipt["status"] == "confirmed"
    assert receipt["photo_delete_after"] is not None

    from app.cleanup import delete_expired_photos
    from app.db import new_session

    with new_session() as db:
        assert delete_expired_photos(db, datetime.now(timezone.utc) + timedelta(days=9)) == 0
        assert delete_expired_photos(db, datetime.now(timezone.utc) + timedelta(days=11)) == 1

    receipt = client.get(f"/api/receipts/{receipt_id}").json()
    assert receipt["has_photo"] is False and receipt["photo_deleted_at"] is not None
    assert client.get(f"/api/receipts/{receipt_id}/photo").status_code == 410
    assert client.get("/api/stats/month", params={"month": "2026-10"}).json()["total_grosze"] == 4599


def test_ocr_failure_marks_receipt(client, monkeypatch):
    def boom(path):
        raise ollama.OcrError("Ollama niedostępna")

    monkeypatch.setattr(ollama, "read_receipt", boom)
    r = client.post("/api/receipts", files={"file": ("p.jpg", jpeg_bytes(), "image/jpeg")})
    receipt = client.get(f"/api/receipts/{r.json()['id']}").json()
    assert receipt["status"] == "failed" and "niedostępna" in receipt["ocr_error"]


def test_normalize_ocr_output():
    out = ollama.normalize({"merchant": "Orlen", "date": "2026-10-1", "total": "123,45 PLN", "items": [{"name": "PB95", "price": 120}]})
    assert out == {"merchant": "Orlen", "date": None, "total_grosze": 12345, "items": [{"name": "PB95", "price_grosze": 12000}]}


def test_suggests_category_from_previous_merchant(client, monkeypatch):
    fuel = client.get("/api/categories").json()[2]["id"]
    client.post("/api/expenses", json={"amount_grosze": 20000, "spent_on": "2026-10-01", "category_id": fuel, "merchant": "Orlen"})
    monkeypatch.setattr(
        ollama, "read_receipt", lambda path: ({"merchant": "ORLEN", "date": None, "total_grosze": 100, "items": []}, "")
    )
    r = client.post("/api/receipts", files={"file": ("p.jpg", jpeg_bytes(), "image/jpeg")})
    assert client.get(f"/api/receipts/{r.json()['id']}").json()["ocr_result"]["suggested_category_id"] == fuel


def test_serves_mobile_app(client):
    r = client.get("/")
    assert r.status_code == 200 and "PixBill" in r.text
    assert client.get("/manifest.webmanifest").status_code == 200
    assert client.get("/api/health").json() == {"status": "ok"}


def test_recurring_generation(client):
    from datetime import date

    from app.db import new_session
    from app.models import Expense, RecurringExpense
    from app.recurring import generate_recurring

    with new_session() as db:
        item = RecurringExpense(description="Czynsz", amount_grosze=200000, category_id=4, day_of_month=10, start_month="2026-08")
        db.add(item)
        db.commit()
        assert generate_recurring(db, date(2026, 10, 5)) == 2
        assert generate_recurring(db, date(2026, 10, 5)) == 0
        assert generate_recurring(db, date(2026, 10, 10)) == 1
        days = sorted(e.spent_on for e in db.query(Expense).filter_by(recurring_id=item.id))
        assert days == [date(2026, 8, 10), date(2026, 9, 10), date(2026, 10, 10)]
        item.active = False
        db.commit()
        assert generate_recurring(db, date(2027, 1, 20)) == 0


def test_recurring_api(client):
    from datetime import date

    cat = first_category(client)
    today = date.today()
    r = client.post("/api/recurring", json={"description": "Internet", "amount_grosze": 6000, "category_id": cat, "day_of_month": 1})
    assert r.status_code == 201
    rid = r.json()["id"]
    expenses = client.get("/api/expenses", params={"month": today.strftime("%Y-%m")}).json()
    assert [e["source"] for e in expenses if e["recurring_id"] == rid] == ["recurring"]

    assert client.patch(f"/api/recurring/{rid}", json={"amount_grosze": 7000, "active": False}).json()["active"] is False
    assert client.post("/api/recurring", json={"description": "X", "amount_grosze": 1, "category_id": cat, "day_of_month": 31}).status_code == 422
    assert client.delete(f"/api/recurring/{rid}").status_code == 204
    assert len(client.get("/api/expenses", params={"month": today.strftime("%Y-%m")}).json()) == 1


def test_offline_resend_is_idempotent(client):
    cat = first_category(client)
    body = {"amount_grosze": 1234, "spent_on": "2026-10-10", "category_id": cat, "client_id": "tel-abc-1"}
    first = client.post("/api/expenses", json=body).json()
    second = client.post("/api/expenses", json=body).json()
    assert first["id"] == second["id"]
    assert len(client.get("/api/expenses", params={"month": "2026-10"}).json()) == 1


def test_receipt_resend_is_idempotent(client, monkeypatch):
    monkeypatch.setattr(ollama, "read_receipt", lambda path: ({"merchant": None, "date": None, "total_grosze": None, "items": []}, ""))
    send = lambda: client.post(
        "/api/receipts", files={"file": ("p.jpg", jpeg_bytes(), "image/jpeg")}, data={"client_id": "tel-rec-1"}
    ).json()["id"]
    assert send() == send()
    assert len(client.get("/api/receipts").json()) == 1


def test_receipt_items_search_and_products(client, monkeypatch):
    food = first_category(client)
    for day, butter in [("2026-09-03", 699), ("2026-10-04", 749)]:
        r = client.post(
            "/api/expenses",
            json={
                "amount_grosze": butter + 1200, "spent_on": day, "category_id": food, "merchant": "Żabka",
                "items": [{"name": "Masło Extra A", "price_grosze": butter}, {"name": "Kawa ziarnista", "price_grosze": 1500},
                          {"name": "Rabat", "price_grosze": -300}],
            },
        )
        assert r.status_code == 201
    expense = r.json()
    assert [i["name"] for i in expense["items"]] == ["Masło Extra A", "Kawa ziarnista", "Rabat"]
    assert expense["items"][0]["product_key"] == "maslo extra"

    assert len(client.get("/api/expenses/search", params={"q": "zabka"}).json()) == 2
    assert len(client.get("/api/expenses/search", params={"q": "MASLO", "date_from": "2026-10-01"}).json()) == 1
    assert client.get("/api/expenses/search", params={"q": "herbata"}).json() == []

    products = client.get("/api/products").json()
    butter = next(p for p in products if p["product_key"] == "maslo extra")
    assert butter["count"] == 2 and butter["total_grosze"] == 1448 and butter["last_price_grosze"] == 749
    assert all(p["name"] != "Rabat" for p in products)
    assert len(client.get("/api/products", params={"month": "2026-10", "q": "kawa"}).json()) == 1
    history = client.get("/api/products/history", params={"key": "maslo extra"}).json()
    assert [h["price_grosze"] for h in history] == [699, 749]

    r = client.patch(f"/api/expenses/{expense['id']}", json={"items": [{"name": "Herbata", "price_grosze": 900}]})
    assert [i["name"] for i in r.json()["items"]] == ["Herbata"]
    assert client.patch(f"/api/expenses/{expense['id']}", json={"description": "x"}).json()["items"][0]["name"] == "Herbata"


def test_confirm_receipt_saves_items(client, monkeypatch):
    monkeypatch.setattr(ollama, "read_receipt", lambda path: ({"merchant": "Lidl", "date": None, "total_grosze": 500, "items": []}, ""))
    receipt_id = client.post("/api/receipts", files={"file": ("p.jpg", jpeg_bytes(), "image/jpeg")}).json()["id"]
    r = client.post(
        f"/api/receipts/{receipt_id}/confirm",
        json={"amount_grosze": 500, "spent_on": "2026-10-20", "category_id": first_category(client),
              "items": [{"name": "Bułka", "price_grosze": 500}]},
    )
    assert r.status_code == 201 and r.json()["items"][0]["name"] == "Bułka"


def test_negative_item_price_parsed():
    assert ollama.parse_amount("-3,00") == -300
    assert ollama.normalize({"items": [{"name": "Rabat", "price": "-1.50"}]})["items"][0]["price_grosze"] == -150


def test_export_xlsx(client):
    from openpyxl import load_workbook

    client.post(
        "/api/expenses",
        json={"amount_grosze": 1234, "spent_on": "2026-10-02", "category_id": first_category(client),
              "description": "zakupy", "items": [{"name": "Chleb", "price_grosze": 1234}]},
    )
    r = client.get("/api/export/xlsx", params={"date_from": "2026-10-01"})
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Wydatki", "Pozycje"]
    assert wb["Wydatki"]["B2"].value == 12.34
    assert wb["Pozycje"]["B2"].value == "Chleb"


def test_deposit_receipt_keeps_photo(client, monkeypatch):
    from app.cleanup import delete_expired_photos
    from app.db import new_session

    monkeypatch.setattr(ollama, "read_receipt", lambda path: ({"merchant": "Lidl", "date": None, "total_grosze": 1000, "items": []}, ""))
    ids = [client.post("/api/receipts", files={"file": ("p.jpg", jpeg_bytes(), "image/jpeg")}).json()["id"] for _ in range(2)]
    for rid, keep in zip(ids, [True, False]):
        r = client.post(
            f"/api/receipts/{rid}/confirm",
            json={"amount_grosze": 1000, "spent_on": "2026-10-20", "category_id": first_category(client), "keep_photo": keep,
                  "items": [{"name": "Woda 1,5L", "price_grosze": 500}, {"name": "Kaucja butelka", "price_grosze": 50}]},
        )
        assert r.status_code == 201

    kept = client.get(f"/api/receipts/{ids[0]}").json()
    assert kept["keep_photo"] and kept["photo_delete_after"] is None and kept["deposit_grosze"] == 50
    with new_session() as db:
        assert delete_expired_photos(db, now=datetime.now(timezone.utc) + timedelta(days=365)) == 1
    assert client.get(f"/api/receipts/{ids[0]}/photo").status_code == 200
    assert client.get(f"/api/receipts/{ids[1]}/photo").status_code == 410
    assert [r["id"] for r in client.get("/api/receipts", params={"keep_photo": True}).json()] == [ids[0]]

    assert client.patch(f"/api/receipts/{ids[1]}", json={"keep_photo": True}).status_code == 409
    r = client.delete(f"/api/receipts/{ids[0]}/photo").json()
    assert not r["keep_photo"] and not r["has_photo"]
    assert client.get(f"/api/receipts/{ids[0]}/photo").status_code == 410


def test_month_report(client):
    food, fuel = first_category(client), client.get("/api/categories").json()[2]["id"]
    for day, amount, cat, items in [
        ("2026-09-05", 1000, food, [{"name": "Masło", "price_grosze": 699}, {"name": "Kawa", "price_grosze": 2000}]),
        ("2026-10-05", 3000, food, [{"name": "Masło", "price_grosze": 799}, {"name": "Kawa", "price_grosze": 1800}]),
        ("2026-10-07", 20000, fuel, []),
    ]:
        client.post("/api/expenses", json={"amount_grosze": amount, "spent_on": day, "category_id": cat, "description": "x", "items": items})
    r = client.get("/api/stats/report", params={"month": "2026-10"}).json()
    assert r["total_grosze"] == 23000 and r["expense_count"] == 2
    assert r["top_category"]["category_id"] == fuel
    assert r["biggest_category_increase"]["category_id"] == food
    assert r["biggest_expense"]["amount_grosze"] == 20000
    assert [(p["name"], p["previous_grosze"], p["current_grosze"]) for p in r["price_rises"]] == [("Masło", 699, 799)]
    assert [p["name"] for p in r["price_drops"]] == ["Kawa"]
